"""Text extraction for PDF / DOCX / TXT / image uploads.

The pipeline prefers native (digital) text and only falls back to OCR when a
page carries no usable text layer — that keeps scanned circulars and Malayalam
hard-copy scans readable without paying the OCR cost for every document.
"""

from __future__ import annotations

import logging
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from ..config import settings

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".log", ".json"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"}
DOC_EXTENSIONS = {".docx", ".doc"}
PDF_EXTENSIONS = {".pdf"}

# A page with fewer usable characters than this is treated as scanned.
MIN_NATIVE_CHARS_PER_PAGE = 40

MALAYALAM_RANGE = re.compile(r"[ഀ-ൿ]")
LATIN_RANGE = re.compile(r"[A-Za-z]")


@dataclass
class PageResult:
    page_number: int
    text: str = ""
    ocr_used: bool = False
    ocr_confidence: float | None = None

    @property
    def char_count(self) -> int:
        return len(self.text)


@dataclass
class ExtractionResult:
    text: str = ""
    pages: list[PageResult] = field(default_factory=list)
    page_count: int | None = None
    method: str = ""
    ocr_used: bool = False
    ocr_quality: float | None = None
    language: str = "Unknown"
    warnings: list[str] = field(default_factory=list)

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    @property
    def char_count(self) -> int:
        return len(self.text)


def classify_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in PDF_EXTENSIONS:
        return "PDF"
    if ext in DOC_EXTENSIONS:
        return "DOCX"
    if ext in IMAGE_EXTENSIONS:
        return "IMAGE"
    return "TXT"


def is_supported(filename: str) -> bool:
    ext = Path(filename).suffix.lower()
    return ext in (PDF_EXTENSIONS | DOC_EXTENSIONS | IMAGE_EXTENSIONS | TEXT_EXTENSIONS)


def ocr_available() -> bool:
    if not settings.kmrl_ocr_enabled:
        return False
    if shutil.which("tesseract") is None:
        return False
    try:
        import pytesseract  # noqa: F401
    except ImportError:  # pragma: no cover - depends on environment
        return False
    return True


def detect_language(text: str) -> str:
    sample = text[:20_000]
    malayalam = len(MALAYALAM_RANGE.findall(sample))
    latin = len(LATIN_RANGE.findall(sample))
    if malayalam == 0 and latin == 0:
        return "Unknown"
    if malayalam == 0:
        return "English"
    if latin == 0:
        return "Malayalam"
    ratio = malayalam / (malayalam + latin)
    if ratio > 0.75:
        return "Malayalam"
    if ratio > 0.08:
        return "Bilingual (EN/ML)"
    return "English"


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- OCR
def _ocr_image(image) -> tuple[str, float | None]:
    """Run tesseract on a PIL image, returning (text, mean confidence 0-100)."""
    import pytesseract

    languages = settings.kmrl_ocr_languages
    try:
        data = pytesseract.image_to_data(
            image, lang=languages, output_type=pytesseract.Output.DICT
        )
    except pytesseract.TesseractError:
        # Requested language pack unavailable — retry with English only.
        data = pytesseract.image_to_data(
            image, lang="eng", output_type=pytesseract.Output.DICT
        )

    words: list[str] = []
    confidences: list[float] = []
    for word, conf in zip(data.get("text", []), data.get("conf", [])):
        token = (word or "").strip()
        if not token:
            continue
        words.append(token)
        try:
            value = float(conf)
        except (TypeError, ValueError):
            continue
        if value >= 0:
            confidences.append(value)

    text = " ".join(words)
    confidence = round(sum(confidences) / len(confidences), 1) if confidences else None
    return text, confidence


def _render_pdf_page(path: Path, page_number: int):
    """Render a single PDF page to a PIL image via poppler."""
    from pdf2image import convert_from_path

    images = convert_from_path(
        str(path),
        dpi=settings.kmrl_ocr_dpi,
        first_page=page_number,
        last_page=page_number,
    )
    return images[0] if images else None


# --------------------------------------------------------------------- extractors
def _extract_pdf(path: Path) -> ExtractionResult:
    from pypdf import PdfReader

    result = ExtractionResult(method="pdf-text")
    try:
        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:  # pragma: no cover - depends on file
                result.warnings.append("PDF is password protected; text may be incomplete.")
        raw_pages = reader.pages
    except Exception as exc:
        raise ValueError(f"Unable to read PDF: {exc}") from exc

    result.page_count = len(raw_pages)
    use_ocr = ocr_available()
    ocr_budget = settings.kmrl_ocr_max_pages
    ocr_confidences: list[float] = []

    for index, page in enumerate(raw_pages, start=1):
        try:
            native = clean_text(page.extract_text() or "")
        except Exception:  # pragma: no cover - malformed page
            native = ""

        page_result = PageResult(page_number=index, text=native)

        if len(native) < MIN_NATIVE_CHARS_PER_PAGE and use_ocr and ocr_budget > 0:
            try:
                image = _render_pdf_page(path, index)
                if image is not None:
                    ocr_text, confidence = _ocr_image(image)
                    ocr_text = clean_text(ocr_text)
                    if len(ocr_text) > len(native):
                        page_result.text = ocr_text
                        page_result.ocr_used = True
                        page_result.ocr_confidence = confidence
                        if confidence is not None:
                            ocr_confidences.append(confidence)
                    ocr_budget -= 1
            except Exception as exc:  # pragma: no cover - depends on poppler
                logger.warning("OCR failed for %s page %s: %s", path.name, index, exc)
                result.warnings.append(f"OCR failed on page {index}.")

        result.pages.append(page_result)

    scanned_pages = sum(1 for page in result.pages if page.ocr_used)
    if scanned_pages:
        result.ocr_used = True
        result.method = "pdf-text+ocr" if scanned_pages < len(result.pages) else "ocr"
        if ocr_confidences:
            result.ocr_quality = round(sum(ocr_confidences) / len(ocr_confidences), 1)

    result.text = clean_text("\n\n".join(p.text for p in result.pages if p.text))

    if not result.text:
        if use_ocr:
            result.warnings.append(
                "No text could be extracted from this PDF, even with OCR."
            )
        else:
            result.warnings.append(
                "This PDF has no text layer and OCR is unavailable on this server."
            )
    return result


def _extract_docx(path: Path) -> ExtractionResult:
    from docx import Document as DocxDocument

    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise ValueError(f"Unable to read Word document: {exc}") from exc

    blocks: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    text = clean_text("\n".join(blocks))
    # Word has no fixed pagination outside a renderer; approximate for display
    # using a conventional 500 words per page.
    approx_pages = max(1, round(len(text.split()) / 500)) if text else None
    pages = [PageResult(page_number=1, text=text)] if text else []
    return ExtractionResult(
        text=text,
        pages=pages,
        page_count=approx_pages,
        method="docx",
    )


def _extract_text_file(path: Path) -> ExtractionResult:
    raw = path.read_text(encoding="utf-8", errors="replace")
    text = clean_text(raw)
    pages = [PageResult(page_number=1, text=text)] if text else []
    return ExtractionResult(text=text, pages=pages, page_count=1, method="plain-text")


def _extract_image(path: Path) -> ExtractionResult:
    result = ExtractionResult(method="ocr", page_count=1)
    if not ocr_available():
        result.warnings.append("OCR is unavailable on this server; image text was not read.")
        return result

    from PIL import Image

    with Image.open(path) as image:
        text, confidence = _ocr_image(image)

    text = clean_text(text)
    result.text = text
    result.ocr_used = True
    result.ocr_quality = confidence
    result.pages = [
        PageResult(page_number=1, text=text, ocr_used=True, ocr_confidence=confidence)
    ]
    if not text:
        result.warnings.append("No readable text was found in this image.")
    return result


def extract_document(path: Path, filename: str | None = None) -> ExtractionResult:
    """Extract text from ``path``, dispatching on the file extension."""
    name = filename or path.name
    ext = Path(name).suffix.lower()

    if ext in PDF_EXTENSIONS:
        result = _extract_pdf(path)
    elif ext in DOC_EXTENSIONS:
        result = _extract_docx(path)
    elif ext in IMAGE_EXTENSIONS:
        result = _extract_image(path)
    elif ext in TEXT_EXTENSIONS:
        result = _extract_text_file(path)
    else:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")

    result.language = detect_language(result.text)
    return result


def poppler_available() -> bool:
    return shutil.which("pdftoppm") is not None


def ocr_engine_version() -> str | None:  # pragma: no cover - environment probe
    if shutil.which("tesseract") is None:
        return None
    try:
        output = subprocess.run(
            ["tesseract", "--version"], capture_output=True, text=True, timeout=5
        )
        return output.stdout.splitlines()[0].strip() if output.stdout else None
    except Exception:
        return None
